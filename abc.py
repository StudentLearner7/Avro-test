import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# Define the input and output files
input_csv = 'input.csv'
output_excel = 'output.xlsx'

def process_csv_to_excel(input_csv, output_excel):
    # Read the data
    data = pd.read_csv(input_csv, skiprows=1, header=None)

    # Delete the 2nd column
    data.drop(data.columns[1], axis=1, inplace=True)
    # Define the new header
    new_header = ['URL', 'pyActivity', 'PreActivity', 'ResponseTime', '#DBCalls', '#API Calls']
    data.columns = new_header[:len(data.columns)]

    # Find the index to stop processing
    stop_index = None
    for i, row in data.iterrows():
        if 'JSP/Servlet' in row['URL']:
            # Remove commas and convert to int
            value = int(row['pyActivity'].replace(',', ''))
            if value < 5000:
                stop_index = i
                break

    # Filter the dataframe if a stop index was found
    if stop_index is not None:
        data = data.iloc[:stop_index]

    # Write to Excel
    wb = Workbook()
    ws = wb.active
    for r in dataframe_to_rows(data, index=False, header=True):
        ws.append(r)

    # Apply grouping
    event_indices = [i for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2) if row[0] == 'Event']
    for start, end in zip(event_indices, event_indices[1:] + [ws.max_row + 1]):
        ws.row_dimensions.group(start+1, end-1, hidden=True)

    wb.save(output_excel)

process_csv_to_excel(input_csv, output_excel)
